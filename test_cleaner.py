import csv
import io
import unittest
from pathlib import Path
from unittest.mock import patch

import app as app_module
from openpyxl import Workbook, load_workbook


class CleanerTests(unittest.TestCase):
    def setUp(self):
        import tempfile

        self.temp_results = tempfile.TemporaryDirectory()
        self.result_dir_patch = patch.object(
            app_module, "RESULT_DIR", Path(self.temp_results.name)
        )
        self.result_dir_patch.start()

        app_module.app.config["TESTING"] = True
        self.client = app_module.app.test_client()
        app_module.app.config["OWNER_USERNAME"] = "test-owner"
        from werkzeug.security import generate_password_hash
        app_module.app.config["OWNER_PASSWORD_HASH"] = generate_password_hash(
            "test-only-password"
        )
        with self.client.session_transaction() as sess:
            sess["owner_authenticated"] = True
            sess["owner_username"] = "test-owner"
            sess["csrf_token"] = "test-csrf-token"
        self.created_files = []

    def tearDown(self):
        for path in self.created_files:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass

        self.result_dir_patch.stop()
        self.temp_results.cleanup()

    def track_download(self, response):
        data = response.get_json()
        if data and data.get("output"):
            self.created_files.append(
                app_module.RESULT_DIR / data["output"]
            )
        return data

    def test_csv_removes_blank_and_duplicate_rows(self):
        content = (
            "name,city\n"
            "Naim,Dhaka\n"
            "\n"
            "Naim,Dhaka\n"
            "Rafi,Khulna\n"
        ).encode("utf-8")

        response = self.client.post(
            "/api/clean",
            data={"file": (io.BytesIO(content), "sample.csv")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        data = self.track_download(response)
        self.assertEqual(data["rows_kept"], 3)
        self.assertEqual(data["duplicate_rows_removed"], 1)
        self.assertEqual(data["blank_rows_removed"], 1)

        download = self.client.get(data["download_url"])
        self.assertEqual(download.status_code, 200)

        try:
            rows = list(csv.reader(
                io.StringIO(download.data.decode("utf-8-sig"))
            ))
        finally:
            download.close()
        self.assertEqual(
            rows,
            [["name", "city"], ["Naim", "Dhaka"], ["Rafi", "Khulna"]],
        )

    def test_rejects_missing_file(self):
        response = self.client.post("/api/clean", data={})
        self.assertEqual(response.status_code, 400)

    def test_rejects_unsupported_extension(self):
        response = self.client.post(
            "/api/clean",
            data={"file": (io.BytesIO(b"hello"), "sample.txt")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 400)

    def test_excel_preview_streams_rows_and_reports_counts(self):
        source = io.BytesIO()
        workbook = Workbook()
        first = workbook.active
        first.title = "Customers"
        first.append(["name", "score"])
        first.append(["Naim", 10])
        first.append(["Naim", 10])
        first.append([None, None])
        first.append(["  Naim  ", 10])

        second = workbook.create_sheet("Orders")
        second.append(["item", "quantity"])
        second.append(["Book", 3])
        second.append(["Book", 3])

        workbook.save(source)
        workbook.close()
        source.seek(0)

        response = self.client.post(
            "/api/preview",
            data={"file": (source, "preview.xlsx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        data = response.get_json()
        self.assertEqual(data["file_type"], "xlsx")
        self.assertEqual(data["rows_total"], 8)
        self.assertEqual(data["rows_kept"], 4)
        self.assertEqual(data["blank_rows_removed"], 1)
        self.assertEqual(data["duplicate_rows_removed"], 3)
        self.assertEqual(
            [sheet["name"] for sheet in data["sheets"]],
            ["Customers", "Orders"],
        )
        self.assertEqual(
            data["sheets"][0]["sample_rows"],
            [["name", "score"], ["Naim", 10]],
        )
        self.assertEqual(data["sheets"][0]["rows_total"], 5)
        self.assertEqual(data["sheets"][1]["rows_kept"], 2)

    def test_excel_cleaning_and_download(self):
        source = io.BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["name", "score"])
        sheet.append(["Naim", 10])
        sheet.append(["Naim", 10])
        sheet.append([None, None])
        workbook.save(source)
        workbook.close()
        source.seek(0)

        response = self.client.post(
            "/api/clean",
            data={"file": (source, "sample.xlsx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        data = self.track_download(response)
        self.assertEqual(data["rows_kept"], 2)
        self.assertEqual(data["duplicate_rows_removed"], 1)
        self.assertEqual(data["blank_rows_removed"], 1)

        download = self.client.get(data["download_url"])
        self.assertEqual(download.status_code, 200)

        download_bytes = download.data
        download.close()
        cleaned_book = load_workbook(
            io.BytesIO(download_bytes), read_only=True, data_only=True
        )
        try:
            rows = list(cleaned_book.active.iter_rows(values_only=True))
            self.assertEqual(rows, [("name", "score"), ("Naim", 10)])
        finally:
            cleaned_book.close()

    def test_excel_cleans_all_worksheets_and_preserves_numbers(self):
        source = io.BytesIO()
        workbook = Workbook()
        first = workbook.active
        first.title = "Customers"
        first.append(["name", "score"])
        first.append(["Naim", 10])
        first.append(["Naim", 10])

        second = workbook.create_sheet("Orders")
        second.append(["item", "quantity"])
        second.append(["Book", 3])
        second.append(["Book", 3])

        workbook.save(source)
        workbook.close()
        source.seek(0)

        response = self.client.post(
            "/api/clean",
            data={"file": (source, "multi_sheet.xlsx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        data = self.track_download(response)
        self.assertEqual(data["rows_kept"], 4)
        self.assertEqual(data["duplicate_rows_removed"], 2)

        download = self.client.get(data["download_url"])
        self.assertEqual(download.status_code, 200)
        download_bytes = download.data
        download.close()

        cleaned_book = load_workbook(
            io.BytesIO(download_bytes), read_only=True, data_only=True
        )
        try:
            self.assertEqual(cleaned_book.sheetnames, ["Customers", "Orders"])
            self.assertEqual(
                list(cleaned_book["Customers"].iter_rows(values_only=True)),
                [("name", "score"), ("Naim", 10)],
            )
            self.assertEqual(
                list(cleaned_book["Orders"].iter_rows(values_only=True)),
                [("item", "quantity"), ("Book", 3)],
            )
        finally:
            cleaned_book.close()

    def test_download_missing_file_returns_404(self):
        response = self.client.get("/download/not-a-real-file.csv")
        self.assertEqual(response.status_code, 404)


    def test_preview_rejects_invalid_xlsx(self):
        response = self.client.post(
            "/api/preview",
            data={"file": (io.BytesIO(b"not a zip file"), "broken.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("error", response.get_json())

    def test_clean_rejects_invalid_xlsx(self):
        response = self.client.post(
            "/api/clean",
            data={"file": (io.BytesIO(b"not a zip file"), "broken.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("error", response.get_json())


if __name__ == "__main__":
    unittest.main(verbosity=2)
