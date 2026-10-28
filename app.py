import csv
import io
import json
import urllib.error
import urllib.request
import time
import os
import uuid
from pathlib import Path

from dotenv import load_dotenv
load_dotenv(Path(__file__).resolve().parent / ".env")

import hmac
import secrets
from datetime import timedelta
from flask import Flask, jsonify, render_template, request, send_from_directory, redirect, url_for, session
from werkzeug.security import check_password_hash
from werkzeug.utils import secure_filename
from openpyxl import load_workbook, Workbook

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = Path(
    os.environ.get("UPLOAD_DIR", str(BASE_DIR / "uploads"))
).expanduser().resolve()
RESULT_DIR = Path(
    os.environ.get("RESULT_DIR", str(BASE_DIR / "results"))
).expanduser().resolve()

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
RESULT_DIR.mkdir(parents=True, exist_ok=True)

from automation_engine import automation_bp
from project_intelligence import project_intelligence_bp
from security_detector import detect_security_issues

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024
if os.environ.get("RENDER") and any(not os.environ.get(name, "").strip() for name in ("FLASK_SECRET_KEY", "OWNER_USERNAME", "OWNER_PASSWORD_HASH")):
    raise RuntimeError(
        "FLASK_SECRET_KEY, OWNER_USERNAME and OWNER_PASSWORD_HASH must be set in the Render environment"
    )

app.register_blueprint(automation_bp)
app.register_blueprint(project_intelligence_bp)

app.config["SECRET_KEY"] = os.environ.get(
    "FLASK_SECRET_KEY", "local-development-only-change-before-deployment"
)
app.config["OWNER_USERNAME"] = os.environ.get("OWNER_USERNAME", "").strip()
app.config["OWNER_PASSWORD_HASH"] = os.environ.get("OWNER_PASSWORD_HASH", "").strip()
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = bool(os.environ.get("RENDER"))
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(hours=8)

ALLOWED_EXTENSIONS = {".csv", ".xlsx"}


import zipfile

MAX_XLSX_ZIP_ENTRIES = 4096
MAX_XLSX_ENTRY_UNCOMPRESSED_SIZE = 50 * 1024 * 1024
MAX_XLSX_TOTAL_UNCOMPRESSED_SIZE = 100 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 1000


def validate_xlsx_zip(file):
    stream = file.stream
    try:
        try:
            archive_context = zipfile.ZipFile(stream, "r")
        except zipfile.BadZipFile as exc:
            raise ValueError("Invalid XLSX ZIP file.") from exc

        with archive_context as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ZIP_ENTRIES:
                raise ValueError("Too many XLSX ZIP entries.")

            total = 0
            names = set()
            for entry in entries:
                names.add(entry.filename)
                if entry.file_size > MAX_XLSX_ENTRY_UNCOMPRESSED_SIZE:
                    raise ValueError("XLSX entry exceeds size limit.")
                total += entry.file_size
                if total > MAX_XLSX_TOTAL_UNCOMPRESSED_SIZE:
                    raise ValueError("XLSX total size limit exceeded.")
                ratio = entry.file_size / max(entry.compress_size, 1)
                if entry.file_size and ratio > MAX_XLSX_COMPRESSION_RATIO:
                    raise ValueError("XLSX compression ratio too high.")

            required = {"[Content_Types].xml", "xl/workbook.xml"}
            if not required.issubset(names):
                raise ValueError("Required XLSX ZIP entries are missing.")
    finally:
        stream.seek(0)



def clean_rows(rows):
    """Remove blank and duplicate rows while preserving non-text values."""
    cleaned = []
    seen = set()
    blank_removed = 0
    duplicate_removed = 0

    for row in rows:
        values = [
            value.strip() if isinstance(value, str) else value
            for value in row
        ]

        if all(value is None or value == "" for value in values):
            blank_removed += 1
            continue

        key = tuple(
            "" if value is None else value
            for value in values
        )

        if key in seen:
            duplicate_removed += 1
            continue

        seen.add(key)
        cleaned.append(values)

    return cleaned, blank_removed, duplicate_removed



def owner_auth_configured():
    return bool(
        app.config.get("OWNER_USERNAME")
        and app.config.get("OWNER_PASSWORD_HASH")
        and app.config.get("SECRET_KEY")
    )


@app.before_request
def require_owner_login():
    if request.endpoint in {"login", "logout", "static"}:
        return None

    if not owner_auth_configured():
        if request.path.startswith("/api/") or request.path.startswith("/download/"):
            return jsonify({
                "status": "configuration_error",
                "error": "Owner login configuration is missing."
            }), 503
        return redirect(url_for("login"))

    if session.get("owner_authenticated") is True:
        return None

    if request.path.startswith("/api/") or request.path.startswith("/download/"):
        return jsonify({
            "status": "unauthorized",
            "error": "লগইন প্রয়োজন।"
        }), 401

    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if not owner_auth_configured():
        return render_template(
            "login.html",
            csrf_token="",
            error="লগইন কনফিগারেশন তৈরি হয়নি। মালিককে সার্ভার সেটআপ করতে হবে।"
        ), 503

    if session.get("owner_authenticated") is True:
        return redirect(url_for("home"))

    token = session.setdefault("csrf_token", secrets.token_urlsafe(32))

    if request.method == "GET":
        return render_template("login.html", csrf_token=token, error=None)

    supplied_token = request.form.get("csrf_token", "")
    if not supplied_token or not hmac.compare_digest(str(token), supplied_token):
        return render_template(
            "login.html",
            csrf_token=token,
            error="অনুরোধ যাচাই করা যায়নি। পেজটি আবার খুলুন।"
        ), 400

    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    expected_username = app.config["OWNER_USERNAME"]
    password_hash = app.config["OWNER_PASSWORD_HASH"]

    try:
        valid_password = check_password_hash(password_hash, password)
    except (ValueError, TypeError):
        valid_password = False

    if not hmac.compare_digest(username, expected_username) or not valid_password:
        return render_template(
            "login.html",
            csrf_token=token,
            error="ব্যবহারকারীর নাম অথবা পাসওয়ার্ড সঠিক নয়।"
        ), 401

    session.clear()
    session.permanent = True
    session["owner_authenticated"] = True
    session["owner_username"] = expected_username
    session["csrf_token"] = secrets.token_urlsafe(32)
    return redirect(url_for("home"))


@app.post("/logout")
def logout():
    expected_token = session.get("csrf_token", "")
    supplied_token = request.form.get("csrf_token", "")

    if (
        not expected_token
        or not supplied_token
        or not hmac.compare_digest(str(expected_token), supplied_token)
    ):
        return "অনুরোধ যাচাই করা যায়নি।", 400

    session.clear()
    return redirect(url_for("login"))


@app.get("/")
def home():
    return render_template("index.html")




def call_gemini(message):
    api_key = os.environ.get("GEMINI_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("GEMINI_API_KEY is not configured")

    model = os.environ.get("GEMINI_MODEL", "gemini-3.8-flash").strip()
    payload = {
        "system_instruction": {
            "parts": [{
                "text": (
                    "You are Naim AI Work Agent. Reply clearly in the "
                    "user's language, preferably Bengali when appropriate. "
                    "Help with Python, CSV and Excel workflows. Never claim "
                    "to have modified files unless a tool actually did so."
                )
            }]
        },
        "contents": [{
            "role": "user",
            "parts": [{"text": message}]
        }],
        "generationConfig": {"maxOutputTokens": 700}
    }

    url = (
        "https://generativelanguage.googleapis.com/v1beta/models/"
        + model + ":generateContent"
    )
    req = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "x-goog-api-key": api_key,
            "Content-Type": "application/json",
        },
        method="POST",
    )

    with urllib.request.urlopen(req, timeout=45) as response:
        result = json.loads(response.read().decode("utf-8"))

    candidates = result.get("candidates", [])
    parts = (
        candidates[0].get("content", {}).get("parts", [])
        if candidates else []
    )
    answer = "\n".join(
        part["text"] for part in parts
        if isinstance(part, dict)
        and isinstance(part.get("text"), str)
    ).strip()

    if not answer:
        raise ValueError("Gemini returned no text.")
    return answer


def ai_is_configured():
    return any(
        os.environ.get(name, "").strip()
        for name in (
            "GEMINI_API_KEY",
            "GROQ_API_KEY",
            "OPENAI_API_KEY",
        )
    )


def call_ai(message):
    providers = (
        ("GEMINI_API_KEY", call_gemini),
        ("GROQ_API_KEY", call_groq),
        ("OPENAI_API_KEY", call_openai),
    )
    last_error = None

    for key_name, function in providers:
        if not os.environ.get(key_name, "").strip():
            continue

        # Retry once only for temporary server errors or network failures.
        for attempt in range(2):
            try:
                answer = function(message)
                provider = key_name.removesuffix("_API_KEY")
                model_env = {
                    "GEMINI_API_KEY": "GEMINI_MODEL",
                    "GROQ_API_KEY": "GROQ_MODEL",
                    "OPENAI_API_KEY": "OPENAI_MODEL",
                }[key_name]
                default_model = {
                    "GEMINI_API_KEY": "gemini-3.8-flash",
                    "GROQ_API_KEY": "openai/gpt-oss-20b",
                    "OPENAI_API_KEY": "gpt-5-mini",
                }[key_name]
                model = os.environ.get(
                    model_env, default_model
                ).strip() or default_model
                app.logger.info(
                    "AI provider succeeded: provider=%s model=%s",
                    provider, model
                )
                return answer
            except Exception as exc:
                last_error = exc
                status_code = getattr(exc, "code", None)

                if isinstance(exc, urllib.error.HTTPError):
                    try:
                        error_body = exc.read().decode(
                            "utf-8", "replace"
                        )
                        error_json = json.loads(error_body)
                        error_info = error_json.get("error", {})
                        if isinstance(error_info, dict):
                            app.logger.warning(
                                "%s API error details: type=%s code=%s message=%s",
                                key_name.removesuffix("_API_KEY"),
                                error_info.get("type", "unknown"),
                                error_info.get("code", "unknown"),
                                error_info.get("message", "No message"),
                            )
                        else:
                            app.logger.warning(
                                "%s API returned an unrecognized error format",
                                key_name.removesuffix("_API_KEY"),
                            )
                    except Exception:
                        app.logger.warning(
                            "%s API error body could not be parsed",
                            key_name.removesuffix("_API_KEY"),
                        )

                retryable = (
                    status_code in (500, 502, 503, 504)
                    or (
                        status_code is None
                        and isinstance(exc, urllib.error.URLError)
                    )
                    or isinstance(exc, TimeoutError)
                )

                app.logger.warning(
                    "%s provider failed: %s%s (attempt %s/2)",
                    key_name.removesuffix("_API_KEY"),
                    type(exc).__name__,
                    f" (HTTP {status_code})" if status_code else "",
                    attempt + 1,
                )

                if retryable and attempt == 0:
                    time.sleep(1)
                    continue

                break

    if last_error is not None:
        raise last_error

    raise RuntimeError("No AI provider is configured")

def call_groq(message):
    api_key = os.environ.get("GROQ_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("GROQ_API_KEY is not configured")

    model = os.environ.get(
        "GROQ_MODEL", "openai/gpt-oss-20b"
    ).strip()

    payload = {
        "model": model,
        "messages": [
            {
                "role": "system",
                "content": (
                    "You are Naim AI Work Agent. Reply helpfully and clearly "
                    "in the user's language, preferably Bengali when appropriate. "
                    "Help with Python, CSV and Excel workflows. Never claim "
                    "to have modified files unless a tool actually did so."
                ),
            },
            {"role": "user", "content": message},
        ],
        "max_completion_tokens": 700,
    }

    req = urllib.request.Request(
        "https://api.groq.com/openai/v1/chat/completions",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": "Bearer " + api_key,
            "Content-Type": "application/json",
            "Accept": "application/json",
            "User-Agent": "naim-ai-work-agent/1.0",
        },
        method="POST",
    )

    with urllib.request.urlopen(req, timeout=35) as response:
        result = json.loads(response.read().decode("utf-8"))

    choices = result.get("choices", [])
    answer = (
        choices[0].get("message", {}).get("content", "")
        if choices else ""
    )
    if not isinstance(answer, str) or not answer.strip():
        raise ValueError("The AI returned no text.")
    return answer.strip()


def call_openai(message):
    api_key = os.environ.get("OPENAI_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY is not configured")

    model = os.environ.get("OPENAI_MODEL", "gpt-5-mini").strip()
    payload = {
        "model": model,
        "instructions": (
            "You are Naim AI Work Agent. Reply helpfully and clearly in "
            "the user's language, preferably Bengali when appropriate. "
            "You can explain tasks and help with CSV/Excel workflows. "
            "Do not claim to have modified files unless a tool actually did so."
        ),
        "input": message,
        "max_output_tokens": 700
    }

    req = urllib.request.Request(
        "https://api.openai.com/v1/responses",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": "Bearer " + api_key,
            "Content-Type": "application/json",
        },
        method="POST",
    )

    with urllib.request.urlopen(req, timeout=35) as response:
        result = json.loads(response.read().decode("utf-8"))

    pieces = []
    for item in result.get("output", []):
        if item.get("type") == "message":
            for part in item.get("content", []):
                if part.get("type") == "output_text" and part.get("text"):
                    pieces.append(part["text"])

    answer = "\n".join(pieces).strip()
    if not answer:
        raise ValueError("The AI returned no text.")
    return answer


def call_tavily_search(query):
    api_key = os.environ.get("TAVILY_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("TAVILY_API_KEY is not configured")

    payload = {
        "api_key": api_key,
        "query": query,
        "search_depth": "basic",
        "max_results": 5,
        "include_answer": False,
    }
    req = urllib.request.Request(
        "https://api.tavily.com/search",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )

    with urllib.request.urlopen(req, timeout=20) as response:
        result = json.loads(response.read().decode("utf-8"))

    return [
        {
            "title": str(item.get("title", "Search result"))[:200],
            "url": str(item.get("url", "")),
            "content": str(item.get("content", ""))[:1200],
        }
        for item in result.get("results", [])
        if isinstance(item, dict) and item.get("url")
    ]




@app.post("/api/project-intelligence/security-advice")
def security_advice():
    """Return AI advice for a limited, sanitized security report."""
    if not ai_is_configured():
        return jsonify({
            "status": "unavailable",
            "message": "AI পরামর্শের জন্য সার্ভারে AI provider কনফিগার করা নেই।",
        }), 503

    try:
        report = detect_security_issues(BASE_DIR)
        safe_findings = []
        for item in report.get("findings", [])[:30]:
            safe_findings.append({
                "severity": str(item.get("severity", ""))[:30],
                "title": str(item.get("title", ""))[:160],
                "file": str(item.get("file", ""))[:240],
                "line": item.get("line"),
                "detail": str(item.get("detail", ""))[:500],
                "solution": str(item.get("solution", ""))[:500],
                "confidence": str(item.get("confidence", "heuristic"))[:40],
            })

        prompt = (
            "You are a cautious code-security reviewer. Answer in Bengali. "
            "Review only the supplied heuristic findings. Prioritize actionable "
            "risks, explain uncertainty and false positives, and give safe "
            "remediation guidance. Do not claim a vulnerability is confirmed. "
            "Do not ask for or invent secrets. Do not modify files or execute code. "
            "The report contains locations and summaries only, not source code.\n\n"
            "Report summary: "
            + json.dumps(report.get("summary", {}), ensure_ascii=False)
            + "\nFindings: "
            + json.dumps(safe_findings, ensure_ascii=False)
        )

        advice = call_ai(prompt)
        return jsonify({
            "status": "ok",
            "read_only": True,
            "advice": advice,
            "findings_considered": len(safe_findings),
            "notice": "AI পরামর্শ যাচাই করে নিন; এটি স্বয়ংক্রিয় patch নয়।",
        })
    except Exception as exc:
        app.logger.warning(
            "Security advice request failed: %s", type(exc).__name__
        )
        return jsonify({
            "status": "error",
            "message": "AI পরামর্শ তৈরি করা যায়নি। পরে আবার চেষ্টা করুন।",
        }), 502


@app.get("/api/status")
def status():
    configured = ai_is_configured()
    return jsonify({
        "app": "Naim AI Work Agent",
        "status": "ok",
        "mode": "cloud-ai-configured" if configured else "local-tools",
        "ai_configured": configured,
    })


@app.post("/api/chat")
def chat():
    data = request.get_json(silent=True) or {}
    message = str(data.get("message", "")).strip()

    if not message:
        return jsonify({"error": "আপনার কাজের নির্দেশটি লিখুন।"}), 400

    if len(message) > 4000:
        return jsonify({"error": "বার্তাটি ৪০০০ অক্ষরের মধ্যে রাখুন।"}), 400

    search_query = ""
    for prefix in ("সার্চ:", "ওয়েবে খুঁজুন:", "ওয়েবে খুঁজুন:", "search:"):
        if message.lower().startswith(prefix.lower()):
            search_query = message[len(prefix):].strip()
            break

    if search_query:
        if not search_query:
            return jsonify({"error": "সার্চের বিষয় লিখুন।"}), 400

        try:
            results = call_tavily_search(search_query)

            if not results:
                return jsonify({
                    "reply": "এই বিষয়ে কোনো সার্চ ফলাফল পাওয়া যায়নি।",
                    "mode": "web-search",
                })

            context = (
                "Answer the user's question in their language, preferably Bengali. "
                "Use these web search results as evidence. Do not invent facts. "
                "Mention uncertainty when needed and include relevant source URLs.\n\n"
                "Question: " + search_query + "\n\n"
            )
            for index, item in enumerate(results, 1):
                context += (
                    f"Result {index}: {item['title']}\n"
                    f"URL: {item['url']}\n"
                    f"Content: {item['content']}\n\n"
                )

            ai_configured = ai_is_configured()

            if ai_configured:
                try:
                    reply = call_ai(context)
                    return jsonify({"reply": reply, "mode": "web-search-ai"})
                except Exception as exc:
                    app.logger.warning(
                        "AI summary failed during web search: %s",
                        type(exc).__name__,
                    )

            lines = ["ওয়েব সার্চ ফলাফল:"]
            for item in results:
                lines.append(
                    f"\n{item['title']}\n{item['content']}\n{item['url']}"
                )
            return jsonify({
                "reply": "\n".join(lines),
                "mode": "web-search-fallback",
            })

        except urllib.error.HTTPError as exc:
            app.logger.warning("Tavily HTTP error: status=%s", exc.code)
            return jsonify({
                "error": f"Tavily সার্চ ব্যর্থ হয়েছে (HTTP {exc.code})। "
                         "Render-এর TAVILY_API_KEY পরীক্ষা করুন।"
            }), 502
        except (urllib.error.URLError, TimeoutError):
            return jsonify({
                "error": "Tavily সার্চে সংযোগ করা যায়নি। আবার চেষ্টা করুন।"
            }), 502
        except Exception:
            app.logger.exception(
                "Tavily search failed; secret values are not logged."
            )
            return jsonify({
                "error": "Tavily সার্চে সমস্যা হয়েছে। Render-এর লগ পরীক্ষা করুন।"
            }), 502

    if ai_is_configured():
        try:
            reply = call_ai(message)
            return jsonify({"reply": reply, "mode": "cloud-ai"})
        except urllib.error.HTTPError as exc:
            detail = ""
            error_type = ""
            error_code = ""

            raw_body = ""
            try:
                raw_body = exc.read().decode("utf-8", "replace")
                body = json.loads(raw_body)
                error = body.get("error", {})

                if isinstance(error, dict):
                    detail = str(
                        error.get("message")
                        or error.get("detail")
                        or body.get("message")
                        or body.get("detail")
                        or raw_body
                    )
                    error_type = str(error.get("type") or "")
                    error_code = str(error.get("code") or "")
                else:
                    detail = str(
                        error or body.get("message")
                        or body.get("detail") or raw_body
                    )
            except (json.JSONDecodeError, ValueError):
                detail = raw_body or "Provider returned an empty error response."
            except Exception:
                detail = raw_body or "Could not read provider error response."

            for secret_name in ("GEMINI_API_KEY", "GROQ_API_KEY", "OPENAI_API_KEY"):
                secret = os.environ.get(secret_name, "").strip()
                if secret:
                    detail = detail.replace(secret, "[REDACTED]")

            app.logger.warning(
                "AI provider HTTP error: status=%s type=%s code=%s detail=%s",
                exc.code, error_type[:100], error_code[:100], detail[:300]
            )

            if exc.code == 401:
                reply = "AI API Key যাচাই করা যায়নি। Render-এর Environment Variables পরীক্ষা করুন।"
            elif exc.code == 429:
                reply = "AI provider-এর usage limit বা billing সমস্যা হতে পারে। সংশ্লিষ্ট provider-এর quota ও billing পরীক্ষা করুন।"
            elif exc.code == 404:
                reply = "নির্বাচিত AI model পাওয়া যায়নি। GEMINI_MODEL, GROQ_MODEL বা OPENAI_MODEL পরীক্ষা করুন।"
            else:
                reply = "AI API অনুরোধ ব্যর্থ হয়েছে। Render-এর লগ পরীক্ষা করুন।"
            return jsonify({"error": reply}), 502
        except (urllib.error.URLError, TimeoutError):
            return jsonify({
                "error": "AI API-তে সংযোগ করা যায়নি। ইন্টারনেট বা সার্ভারের অবস্থা পরীক্ষা করুন।"
            }), 502
        except (ValueError, json.JSONDecodeError):
            return jsonify({
                "error": "AI থেকে সঠিক উত্তর পাওয়া যায়নি। আবার চেষ্টা করুন।"
            }), 502
        except Exception:
            app.logger.exception("AI provider request failed; secret values are not logged.")
            return jsonify({
                "error": "AI অনুরোধে সমস্যা হয়েছে। Render-এর লগ পরীক্ষা করুন।"
            }), 502

    lower = message.lower()

    if any(word in lower for word in ("help", "সাহায্য", "কি করতে", "কী করতে")):
        reply = (
            "আমি CSV ও XLSX ফাইল থেকে ফাঁকা এবং ডুপ্লিকেট সারি সরাতে পারি। "
            "ক্লাউড AI চালাতে Gemini, Groq অথবা OpenAI API key সেট করতে হবে।"
        )
    elif any(word in lower for word in ("excel", "csv", "ফাইল", "ডেটা", "ডাটা")):
        reply = (
            "আপনার CSV বা XLSX ফাইলটি Upload অংশে দিন। "
            "আমি ফাঁকা ও ডুপ্লিকেট সারি সরিয়ে পরিষ্কার ফাইল তৈরি করব।"
        )
    else:
        reply = (
            "আপনার নির্দেশ পেয়েছি। এখন লোকাল মোড চলছে। "
            "ক্লাউড AI চালাতে Render-এর Environment Variables-এ "
            "GEMINI_API_KEY, GROQ_API_KEY অথবা OPENAI_API_KEY সেট করুন।"
        )

    return jsonify({"reply": reply, "mode": "local-tools"})


@app.post("/api/preview")
def preview_file():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "প্রথমে একটি CSV বা XLSX ফাইল নির্বাচন করুন।"}), 400

    filename = secure_filename(file.filename)
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        return jsonify({"error": "শুধু .csv ও .xlsx ফাইল গ্রহণ করা হয়।"}), 400

    try:
        if extension == ".csv":
            raw = file.read()
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("latin-1")

            rows = list(csv.reader(io.StringIO(text)))
            cleaned, blank, duplicate = clean_rows(rows)
            return jsonify({
                "message": "প্রিভিউ তৈরি হয়েছে। এখনো কোনো ফাইল পরিষ্কার করা হয়নি।",
                "file_type": "csv",
                "rows_total": len(rows),
                "rows_kept": len(cleaned),
                "blank_rows_removed": blank,
                "duplicate_rows_removed": duplicate,
                "sample_rows": cleaned[:6],
                "sheets": [],
            })

        validate_xlsx_zip(file)
        workbook = load_workbook(file, read_only=True, data_only=True)
        sheets = []
        total_rows = total_kept = total_blank = total_duplicate = 0
        try:
            for sheet in workbook.worksheets:
                sheet_rows = sheet_kept = 0
                blank = duplicate = 0
                seen = set()
                sample_rows = []

                for row in sheet.iter_rows(values_only=True):
                    sheet_rows += 1
                    values = [
                        value.strip() if isinstance(value, str) else value
                        for value in row
                    ]

                    if all(value is None or value == "" for value in values):
                        blank += 1
                        continue

                    key = tuple(
                        "" if value is None else value
                        for value in values
                    )
                    if key in seen:
                        duplicate += 1
                        continue

                    seen.add(key)
                    sheet_kept += 1

                    if len(sample_rows) < 6:
                        sample_rows.append([
                            value.isoformat()
                            if hasattr(value, "isoformat")
                            else value
                            for value in values
                        ])

                total_rows += sheet_rows
                total_kept += sheet_kept
                total_blank += blank
                total_duplicate += duplicate

                sheets.append({
                    "name": sheet.title,
                    "rows_total": sheet_rows,
                    "rows_kept": sheet_kept,
                    "blank_rows_removed": blank,
                    "duplicate_rows_removed": duplicate,
                    "sample_rows": sample_rows,
                })
        finally:
            workbook.close()
        return jsonify({
            "message": "প্রিভিউ তৈরি হয়েছে। এখনো কোনো ফাইল পরিষ্কার করা হয়নি।",
            "file_type": "xlsx",
            "rows_total": total_rows,
            "rows_kept": total_kept,
            "blank_rows_removed": total_blank,
            "duplicate_rows_removed": total_duplicate,
            "sample_rows": [],
            "sheets": sheets,
        })
    except Exception:
        app.logger.exception("File preview failed")
        return jsonify({
            "error": "ফাইলটি পড়া যায়নি। বৈধ CSV/XLSX ফাইল কি না পরীক্ষা করুন।"
        }), 400


@app.post("/api/clean")
def clean_file():
    file = request.files.get("file")

    if not file or not file.filename:
        return jsonify({"error": "প্রথমে একটি CSV বা XLSX ফাইল নির্বাচন করুন।"}), 400

    filename = secure_filename(file.filename)
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        return jsonify({"error": "শুধু .csv ও .xlsx ফাইল গ্রহণ করা হয়।"}), 400

    file_id = uuid.uuid4().hex[:12]
    output_name = f"cleaned_{file_id}.csv" if extension == ".csv" else f"cleaned_{file_id}.xlsx"

    try:
        if extension == ".csv":
            raw = file.read()
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("latin-1")

            rows = list(csv.reader(io.StringIO(text)))
            cleaned, blank, duplicate = clean_rows(rows)

            output_path = RESULT_DIR / output_name
            with output_path.open("w", newline="", encoding="utf-8-sig") as output:
                writer = csv.writer(output)
                writer.writerows(cleaned)

        else:
            validate_xlsx_zip(file)
            workbook = load_workbook(file, read_only=True, data_only=True)
            output_book = Workbook(write_only=True)

            rows_kept = 0
            blank = 0
            duplicate = 0

            try:
                for source_sheet in workbook.worksheets:
                    target_sheet = output_book.create_sheet(
                        title=source_sheet.title
                    )
                    seen = set()

                    for row in source_sheet.iter_rows(values_only=True):
                        values = [
                            value.strip() if isinstance(value, str) else value
                            for value in row
                        ]

                        if all(value is None or value == "" for value in values):
                            blank += 1
                            continue

                        key = tuple(
                            "" if value is None else value
                            for value in values
                        )

                        if key in seen:
                            duplicate += 1
                            continue

                        seen.add(key)
                        target_sheet.append(values)
                        rows_kept += 1
            finally:
                workbook.close()

            output_path = RESULT_DIR / output_name
            try:
                output_book.save(output_path)
            finally:
                output_book.close()

    except Exception:
        app.logger.exception("File cleaning failed")
        return jsonify({
            "error": "ফাইলটি পড়া যায়নি। ফাইলটি বৈধ CSV/XLSX কি না পরীক্ষা করুন।"
        }), 400

    return jsonify({
        "message": "ফাইল পরিষ্কার সম্পন্ন হয়েছে।",
        "output": output_name,
        "rows_kept": rows_kept if extension == ".xlsx" else len(cleaned),
        "blank_rows_removed": blank,
        "duplicate_rows_removed": duplicate,
        "download_url": f"/download/{output_name}",
    })


@app.get("/download/<path:filename>")
def download(filename):
    safe_name = secure_filename(filename)
    if not safe_name or not (RESULT_DIR / safe_name).is_file():
        return jsonify({"error": "ফাইল পাওয়া যায়নি।"}), 404

    return send_from_directory(
        RESULT_DIR, safe_name, as_attachment=True
    )


@app.errorhandler(413)
def too_large(_error):
    return jsonify({"error": "ফাইলের সর্বোচ্চ আকার ১০ MB।"}), 413


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="127.0.0.1", port=port, debug=False)
